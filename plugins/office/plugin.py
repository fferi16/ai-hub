"""office – Word/Excel/PowerPoint: megbízható PDF-export, Excel-olvasás és -összehasonlítás, Word-szöveg

A PDF-export külön folyamatban fut (convert.py) saját Office-példánnyal és időkorláttal:
ha elakad (rejtett párbeszédablak), csak az általa indított Office-folyamatokat állítja le.
"""
import json
import re
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

from hubkit import secret, tool  # noqa: F401

HERE = Path(__file__).resolve().parent
PDF_TIMEOUT_S = 100
OFFICE_EXT = {".docx", ".doc", ".rtf", ".odt", ".xlsx", ".xlsm", ".xls", ".csv", ".pptx", ".ppt"}
NOWIN = 0x08000000


def _file(path: str, exts: set[str] | None = None):
    p = Path(path).expanduser()
    if not p.is_file():
        return None, {"ok": False, "reason": f"nincs ilyen fájl: {path}"}
    if exts and p.suffix.lower() not in exts:
        return None, {"ok": False, "reason": f"nem támogatott formátum: {p.suffix}"}
    return p.resolve(), None


def _kill_started(pids_file: Path) -> list[int]:
    try:
        pids = json.loads(pids_file.read_text("utf-8"))
    except (OSError, ValueError):
        return []
    for pid in pids:
        subprocess.run(["taskkill", "/F", "/PID", str(pid)], capture_output=True, creationflags=NOWIN)
    return pids


@tool
def office_to_pdf(path: str, out: str = "", update_toc: bool = True, save_docx: str = "") -> dict:
    """Word/Excel/PowerPoint → PDF (elakadásvédelemmel, ~100 s korlát); Wordnél frissíti a tartalomjegyzéket, save_docx-ba a frissített példányt menti."""
    src, err = _file(path, OFFICE_EXT)
    if err:
        return err
    pdf = Path(out).expanduser() if out else src.with_suffix(".pdf")
    if pdf.suffix.lower() != ".pdf":
        return {"ok": False, "reason": "az out kiterjesztése .pdf legyen"}
    if save_docx and src.suffix.lower() not in {".docx", ".doc", ".rtf", ".odt"}:
        return {"ok": False, "reason": "save_docx csak Word-dokumentumnál használható"}
    pdf.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        job = {"src": str(src), "pdf": str(pdf), "update_toc": update_toc, "docx_out": save_docx or None,
               "pids_file": str(Path(td) / "pids.json"), "result_file": str(Path(td) / "result.json")}
        (Path(td) / "job.json").write_text(json.dumps(job, ensure_ascii=False), "utf-8")
        t0 = time.time()
        proc = subprocess.Popen([sys.executable, str(HERE / "convert.py"), str(Path(td) / "job.json")],
                                stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, creationflags=NOWIN)
        try:
            _, errout = proc.communicate(timeout=PDF_TIMEOUT_S)
        except subprocess.TimeoutExpired:
            proc.kill()
            killed = _kill_started(Path(job["pids_file"]))
            return {"ok": False, "reason": f"időtúllépés ({PDF_TIMEOUT_S} s) – valószínűleg rejtett Office-párbeszéd",
                    "killed_office_pids": killed}
        res_file = Path(job["result_file"])
        if not res_file.exists():
            _kill_started(Path(job["pids_file"]))
            return {"ok": False, "reason": "a konverter összeomlott", "stderr": errout.decode(errors="replace")[-400:]}
        res = json.loads(res_file.read_text("utf-8"))
    if res.get("ok"):
        res.update(kb=round(pdf.stat().st_size / 1024), seconds=round(time.time() - t0, 1))
    return res


# ---------------- Excel ----------------
def _wb(path: str, formulas: bool):
    from openpyxl import load_workbook
    p, err = _file(path, {".xlsx", ".xlsm"})
    if err:
        return None, err
    return load_workbook(p, data_only=not formulas), None


def _val(v):
    if v is None or isinstance(v, (int, float, str, bool)):
        return v
    return str(v)


@tool
def office_xl_sheets(path: str) -> dict:
    """Excel-munkafüzet lapjai: név, használt tartomány, sorok/oszlopok száma."""
    wb, err = _wb(path, False)
    if err:
        return err
    return {"ok": True, "sheets": [{"name": ws.title, "range": ws.dimensions, "rows": ws.max_row, "cols": ws.max_column}
                                   for ws in wb.worksheets]}


@tool
def office_xl_read(path: str, sheet: str = "", cells: str = "", formulas: bool = False, max_cells: int = 2000) -> dict:
    """Excel-tartomány értékei (vagy képletei) soronként; cells pl. 'A1:F40' (üres = használt tartomány); üres cellák nélkül."""
    wb, err = _wb(path, formulas)
    if err:
        return err
    if sheet and sheet not in wb.sheetnames:
        return {"ok": False, "reason": f"nincs ilyen lap: {sheet}", "sheets": wb.sheetnames}
    ws = wb[sheet] if sheet else wb.active
    rng = cells or ws.dimensions
    try:
        grid = ws[rng] if ":" in rng else ((ws[rng],),)
    except ValueError:
        return {"ok": False, "reason": f"hibás tartomány: {rng}"}
    rows, n, cut = {}, 0, False
    for row in grid:
        for c in row:
            if c.value is None:
                continue
            if n >= max_cells:
                cut = True
                break
            rows.setdefault(str(c.row), {})[c.column_letter] = _val(c.value)
            n += 1
        if cut:
            break
    return {"ok": True, "sheet": ws.title, "range": rng, "cells": n, "truncated": cut, "rows": rows}


@tool
def office_xl_diff(path_a: str, path_b: str, sheet: str = "", formulas: bool = False, max_changes: int = 500) -> dict:
    """Két Excel-fájl cellaszintű különbsége (lapszinten is): melyik cella mi volt → mi lett."""
    a, err = _wb(path_a, formulas)
    if err:
        return err
    b, err = _wb(path_b, formulas)
    if err:
        return err
    names = [sheet] if sheet else list(dict.fromkeys(a.sheetnames + b.sheetnames))
    changes, only_a, only_b, cut = [], [], [], False
    for name in names:
        if name not in b.sheetnames:
            only_a.append(name)
            continue
        if name not in a.sheetnames:
            only_b.append(name)
            continue
        wa, wb_ = a[name], b[name]
        for r in range(1, max(wa.max_row, wb_.max_row) + 1):
            for c in range(1, max(wa.max_column, wb_.max_column) + 1):
                va, vb = _val(wa.cell(r, c).value), _val(wb_.cell(r, c).value)
                if va != vb:
                    if len(changes) >= max_changes:
                        cut = True
                        break
                    changes.append({"sheet": name, "cell": wa.cell(r, c).coordinate, "a": va, "b": vb})
            if cut:
                break
    return {"ok": True, "changed": len(changes), "truncated": cut, "only_in_a": only_a, "only_in_b": only_b,
            "changes": changes}


# ---------------- Word ----------------
_ENT = {"&amp;": "&", "&lt;": "<", "&gt;": ">", "&quot;": '"', "&apos;": "'"}


@tool
def office_docx_text(path: str, max_chars: int = 20000) -> dict:
    """Word (.docx) szövege bekezdésenként, a címsorokat '#' jelöléssel."""
    p, err = _file(path, {".docx"})
    if err:
        return err
    with zipfile.ZipFile(p) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    lines = []
    for para in re.findall(r"<w:p[ >].*?</w:p>", xml, flags=re.S):
        text = "".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", para))
        if not text.strip():
            continue
        m = re.search(r'<w:pStyle w:val="(?:Heading|Cmsor)(\d)"', para)
        lines.append(("#" * int(m.group(1)) + " " if m else "") + text)
    text = re.sub("|".join(_ENT), lambda m: _ENT[m.group(0)], "\n".join(lines))
    return {"ok": True, "paragraphs": len(lines), "chars": len(text), "truncated": len(text) > max_chars,
            "text": text[:max_chars]}
